"""Le bug d'origine (bonus) : generate_docx() shellait vers un script
Node.js qui dépendait d'un module npm jamais déclaré nulle part — la
génération échouait toujours sur une installation propre. Ce test
vérifie qu'un vrai fichier .docx est produit, avec le contenu attendu.
"""

from docx import Document
from dpa_generator import DEMO_CONFIG, build_dpa_data, generate_docx


def test_generate_docx_produces_a_real_readable_file(tmp_path):
    data = build_dpa_data(DEMO_CONFIG)
    output = tmp_path / "dpa.docx"

    result_path = generate_docx(data, output)

    assert result_path == output
    assert output.exists()
    assert output.stat().st_size > 1000  # pas un fichier vide/tronqué


def _all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_generated_docx_contains_the_parties_and_clauses(tmp_path):
    data = build_dpa_data(DEMO_CONFIG)
    output = tmp_path / "dpa.docx"
    generate_docx(data, output)

    full_text = _all_text(Document(output))

    assert DEMO_CONFIG["rt_nom"] in full_text
    assert DEMO_CONFIG["st_nom"] in full_text
    for clause in data["clauses"].values():
        assert clause["titre"] in full_text
